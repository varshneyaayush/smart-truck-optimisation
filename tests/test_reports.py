"""
tests/test_reports.py
Tests for in-memory Excel and PDF operational report generation and value consistency.
"""

import io
from openpyxl import load_workbook
import pytest
from core.models import Vehicle, CargoItem
from core.optimizer import PackingOptimizer
from reports.excel_report import generate_excel_report
from reports.pdf_report import generate_pdf_report


@pytest.fixture
def optimization_result():
    veh = Vehicle(
        id="T1",
        label="14 ft Test Truck",
        category="14 ft",
        nominal_volume_m3=15.63,
        max_payload_kg=3500.0,
        internal_length_cm=420.0,
        internal_width_cm=205.0,
        internal_height_cm=190.0,
        door_width_cm=195.0,
        door_height_cm=180.0,
        is_verified=False
    )
    items = [
        CargoItem("CRG-1", "Pallet", 120, 80, 100, 300.0, quantity=2),
        CargoItem("CRG-2", "Fragile Carton", 60, 40, 40, 20.0, quantity=1, fragile=True)
    ]
    opt = PackingOptimizer(veh)
    return opt.pack(items)


def test_excel_report_generation(optimization_result):
    excel_bytes = generate_excel_report(optimization_result)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 2000

    # Parse back with openpyxl to verify sheets and consistency
    wb = load_workbook(io.BytesIO(excel_bytes))
    assert "Executive Summary" in wb.sheetnames
    assert "Placed Cargo Manifest" in wb.sheetnames
    assert "Unplaced Items" in wb.sheetnames
    assert "Validation Findings" in wb.sheetnames

    # Check that loaded weight matches in summary sheet
    ws_summary = wb["Executive Summary"]
    found_weight = False
    for row in ws_summary.iter_rows(values_only=True):
        if row[0] == "Total Loaded Weight":
            assert f"{optimization_result.loaded_weight_kg:,.1f} kg" == row[1]
            found_weight = True
            break
    assert found_weight is True


def test_pdf_report_generation(optimization_result):
    pdf_bytes = generate_pdf_report(optimization_result)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 3000
    # PDF magic bytes
    assert pdf_bytes.startswith(b"%PDF-")
