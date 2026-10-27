"""
reports/excel_report.py
Generates downloadable Excel reports using OpenPyXL.
Contains Summary, Vehicle Details, Placed Items, Unplaced Items, and Validation Findings.
Strictly in-memory generation; does not persist files to disk.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from core.models import OptimizationResult

# Approved brand styling
WINE_RED = "7F011F"
LIGHT_SAND = "F5EBD0"
WARM_GRAY = "EFEBE4"
WHITE = "FFFFFF"
GREEN = "1E7E34"
RED = "C82333"


def _apply_header_style(cell):
    cell.font = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
    cell.fill = PatternFill(start_color=WINE_RED, end_color=WINE_RED, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_sub_header(cell):
    cell.font = Font(name="Segoe UI", size=10, bold=True, color="333333")
    cell.fill = PatternFill(start_color=LIGHT_SAND, end_color=LIGHT_SAND, fill_type="solid")
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _apply_border(cell):
    thin = Side(border_style="thin", color="D3C9B8")
    cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)


def generate_excel_report(result: OptimizationResult) -> bytes:
    """
    Generates a structured, styled Excel workbook in memory.
    """
    wb = Workbook()

    # -------------------------------------------------------------
    # Sheet 1: Executive Summary
    # -------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    # Title block
    ws_summary["A1"] = "OPTIMOVE — SMART TRUCK LOADING OPTIMIZATION REPORT"
    ws_summary["A1"].font = Font(name="Segoe UI", size=15, bold=True, color=WINE_RED)
    ws_summary["A2"] = f"Optimization Status: {result.status} | {result.solver_message}"
    ws_summary["A2"].font = Font(name="Segoe UI", size=10, italic=True)

    summary_rows = [
        ("Vehicle Model / ID", f"{result.vehicle.label} ({result.vehicle.id})"),
        ("Vehicle Verification Status", result.vehicle.verification_status),
        ("Vehicle Dimensions (L x W x H)", f"{result.vehicle.internal_length_cm} x {result.vehicle.internal_width_cm} x {result.vehicle.internal_height_cm} cm"),
        ("Door Opening (W x H)", f"{result.vehicle.door_width_cm} x {result.vehicle.door_height_cm} cm"),
        ("Vehicle Internal Volume", f"{result.vehicle.internal_volume_m3:.2f} m³"),
        ("Maximum Certified Payload", f"{result.vehicle.max_payload_kg:,.1f} kg"),
        ("--------------------------------", "--------------------------------"),
        ("Total Placed Items", f"{len(result.placed_instances)} units"),
        ("Total Unplaced Items", f"{len(result.unplaced_instances)} units"),
        ("Total Loaded Weight", f"{result.loaded_weight_kg:,.1f} kg"),
        ("Remaining Payload Capacity", f"{result.remaining_payload_kg:,.1f} kg"),
        ("Payload Utilization", f"{(result.loaded_weight_kg / result.vehicle.max_payload_kg * 100):.1f}%" if result.vehicle.max_payload_kg > 0 else "N/A"),
        ("Cargo Physical Volume", f"{result.placed_cargo_volume_m3:.3f} m³"),
        ("Volume Utilization", f"{result.volume_utilization_pct:.1f}%"),
        ("Bounding Spatial Utilization", f"{result.spatial_utilization_pct:.1f}%"),
        ("Independent Validation Status", "PASSED (All constraints satisfied)" if result.validation_result and result.validation_result.is_valid else "FAILED / VIOLATIONS RECORDED"),
    ]

    ws_summary.append([])
    ws_summary.append(["METRIC / ATTRIBUTE", "VALUE"])
    _apply_header_style(ws_summary.cell(row=4, column=1))
    _apply_header_style(ws_summary.cell(row=4, column=2))

    for r_idx, (k, v) in enumerate(summary_rows, start=5):
        c1 = ws_summary.cell(row=r_idx, column=1, value=k)
        c2 = ws_summary.cell(row=r_idx, column=2, value=v)
        c1.font = Font(name="Segoe UI", size=10, bold=True if "----------------" not in k else False)
        c2.font = Font(name="Segoe UI", size=10)
        _apply_border(c1)
        _apply_border(c2)

    # -------------------------------------------------------------
    # Sheet 2: Placed Cargo Manifest
    # -------------------------------------------------------------
    ws_placed = wb.create_sheet(title="Placed Cargo Manifest")
    ws_placed.views.sheetView[0].showGridLines = True

    headers_placed = [
        "Seq #", "Instance ID", "Parent Item ID", "Cargo Name",
        "Pos X (cm)", "Pos Y (cm)", "Pos Z (cm)",
        "Length (cm)", "Width (cm)", "Height (cm)",
        "Weight (kg)", "Orientation", "Fragile", "Stackable", "Max Top Load (kg)", "Late Order"
    ]
    ws_placed.append(headers_placed)
    for col_num in range(1, len(headers_placed) + 1):
        _apply_header_style(ws_placed.cell(row=1, column=col_num))

    for p in result.placed_instances:
        row_vals = [
            p.loading_sequence_index or "-",
            p.instance_id,
            p.parent_item_id,
            p.name,
            p.x_cm,
            p.y_cm,
            p.z_cm,
            p.placed_length_cm,
            p.placed_width_cm,
            p.placed_height_cm,
            p.weight_kg,
            p.orientation_label,
            "Yes" if p.fragile else "No",
            "Yes" if p.stackable else "No",
            p.max_top_load_kg,
            "Yes" if p.is_late_order else "No"
        ]
        ws_placed.append(row_vals)
        cur_row = ws_placed.max_row
        for col_num in range(1, len(headers_placed) + 1):
            cell = ws_placed.cell(row=cur_row, column=col_num)
            cell.font = Font(name="Segoe UI", size=9)
            cell.alignment = Alignment(horizontal="center" if col_num not in [2, 4] else "left")
            _apply_border(cell)

    # -------------------------------------------------------------
    # Sheet 3: Unplaced Items
    # -------------------------------------------------------------
    ws_unplaced = wb.create_sheet(title="Unplaced Items")
    ws_unplaced.views.sheetView[0].showGridLines = True

    headers_unplaced = ["Instance ID", "Parent Item ID", "Name", "Length (cm)", "Width (cm)", "Height (cm)", "Weight (kg)", "Verifiable Rejection Reason"]
    ws_unplaced.append(headers_unplaced)
    for col_num in range(1, len(headers_unplaced) + 1):
        _apply_header_style(ws_unplaced.cell(row=1, column=col_num))

    if not result.unplaced_instances:
        ws_unplaced.append(["None", "-", "All cargo items placed successfully", "-", "-", "-", "-", "N/A"])
        for col_num in range(1, len(headers_unplaced) + 1):
            cell = ws_unplaced.cell(row=2, column=col_num)
            cell.font = Font(name="Segoe UI", size=9, italic=True)
            _apply_border(cell)
    else:
        for u_inst, reason in result.unplaced_instances:
            row_vals = [
                u_inst.instance_id,
                u_inst.parent_item_id,
                u_inst.name,
                u_inst.length_cm,
                u_inst.width_cm,
                u_inst.height_cm,
                u_inst.weight_kg,
                reason
            ]
            ws_unplaced.append(row_vals)
            cur_row = ws_unplaced.max_row
            for col_num in range(1, len(headers_unplaced) + 1):
                cell = ws_unplaced.cell(row=cur_row, column=col_num)
                cell.font = Font(name="Segoe UI", size=9)
                cell.alignment = Alignment(horizontal="left" if col_num in [3, 8] else "center")
                _apply_border(cell)

    # -------------------------------------------------------------
    # Sheet 4: Validation Findings
    # -------------------------------------------------------------
    ws_val = wb.create_sheet(title="Validation Findings")
    ws_val.views.sheetView[0].showGridLines = True

    headers_val = ["Constraint Check Description", "Audit Status", "Details / Diagnostics"]
    ws_val.append(headers_val)
    for col_num in range(1, len(headers_val) + 1):
        _apply_header_style(ws_val.cell(row=1, column=col_num))

    if result.validation_result:
        val = result.validation_result
        for chk, ok in val.summary_checks.items():
            readable_chk = chk.replace("_", " ").title()
            ws_val.append([
                readable_chk,
                "PASS" if ok else "FAIL",
                "No violations" if ok else "Violations detected during constraint analysis"
            ])
            r = ws_val.max_row
            ws_val.cell(row=r, column=1).font = Font(name="Segoe UI", size=9)
            stat_cell = ws_val.cell(row=r, column=2)
            stat_cell.font = Font(name="Segoe UI", size=9, bold=True, color=GREEN if ok else RED)
            stat_cell.alignment = Alignment(horizontal="center")
            ws_val.cell(row=r, column=3).font = Font(name="Segoe UI", size=9)
            for c in range(1, 4):
                _apply_border(ws_val.cell(row=r, column=c))

        if val.errors:
            ws_val.append([])
            ws_val.append(["EXPLICIT VIOLATION MESSAGES", "", ""])
            _apply_sub_header(ws_val.cell(row=ws_val.max_row, column=1))
            for err in val.errors:
                ws_val.append([err, "VIOLATION", "Requires adjustment"])
                r = ws_val.max_row
                ws_val.cell(row=r, column=1).font = Font(name="Segoe UI", size=9, color=RED)
                ws_val.cell(row=r, column=2).font = Font(name="Segoe UI", size=9, bold=True, color=RED)
                ws_val.cell(row=r, column=3).font = Font(name="Segoe UI", size=9)

    # Auto-adjust column widths across all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
